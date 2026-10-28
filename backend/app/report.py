"""报告生成：python-docx 输出 .docx。所有结论逐条附证据定位。"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from . import rule_refs

LEVEL_CN = {"high": "高", "medium": "中", "low": "低"}
QC_CN = {"pass": "通过", "fail": "未通过", "insufficient": "材料缺失，无法验证"}


def _set_font(run, size=10.5, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    if color:
        run.font.color.rgb = RGBColor(*color)


def build_report(path: Path, project: dict, docs: list[dict], findings: list[dict],
                 spec: dict, qc_result: dict, missing: list[dict], board_label: str = "资产权属与经营基础",
                 triggers: list[dict] | None = None) -> Path:
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run(f"{project['title']}\n{board_label}核查报告"), 18, True)

    if project.get("is_demo"):
        warn = doc.add_paragraph()
        warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(warn.add_run("【虚构测试样本 · 仅用于系统演示，不得用于任何真实交易或法律结论】"), 11, True, (0xC0, 0x00, 0x00))

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("拟上市主体", project.get("subject_name") or "（未填写）"),
        ("业务板块", board_label),
        ("生成时间", str(project.get("updated_at", "")).replace("T", " ")[:19]),
        ("分析模式", "演示模式（未接模型，规则启发式）" if spec.get("mode") == "mock" else "模型分析（DeepSeek）"),
        ("材料数量", f"{len(docs)} 份"),
        ("质检结论", ("通过" if qc_result.get("passed") else "未通过")
         + f"（核查 {qc_result.get('checked', 0)} 条，阻断项 {qc_result.get('failed', 0)} 条）"),
    ]
    for k, v in rows:
        cells = meta.add_row().cells
        _set_font(cells[0].paragraphs[0].add_run(k), 10.5, True)
        _set_font(cells[1].paragraphs[0].add_run(str(v)), 10.5)

    doc.add_heading("一、核查范围与任务单", level=1)
    doc.add_paragraph(f"核查目标：{spec.get('objective', '')}")
    doc.add_paragraph("核查要点：")
    for cp in spec.get("checks", []):
        doc.add_paragraph(f"{cp['id']} {cp['topic']}：{cp['check_point']}", style="List Bullet")

    doc.add_heading("二、材料清单", level=1)
    if docs:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for i, h in enumerate(["序号", "文件名", "页数", "解析状态"]):
            _set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), 10, True)
        for idx, d in enumerate(docs, start=1):
            cells = table.add_row().cells
            for i, val in enumerate([str(idx), d["filename"], str(d.get("pages", "")), d.get("status", "")]):
                _set_font(cells[i].paragraphs[0].add_run(val), 10)
    else:
        doc.add_paragraph("（无）")

    if missing:
        doc.add_heading("三、缺失材料（未完成核查的前置条件）", level=1)
        doc.add_paragraph("以下材料未在本次上传中提供，相应核查项无法完成：")
        for m in missing:
            doc.add_paragraph(f"{m['material']}——用于 {m['needed_for']}", style="List Bullet")

    doc.add_heading("四、风险清单", level=1)
    if findings:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, h in enumerate(["编号", "等级", "结论", "证据定位", "质检"]):
            _set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), 10, True)
        for f in findings:
            cells = table.add_row().cells
            loc = "；".join(
                f"{c['filename']} p{c['page']}" for c in f["citations"] if c.get("document_id")
            ) or "（材料缺失）"
            values = [f["code"], LEVEL_CN.get(f["risk_level"], f["risk_level"]), f["title"], loc,
                      QC_CN.get(f.get("qc", ""), "")]
            for i, val in enumerate(values):
                _set_font(cells[i].paragraphs[0].add_run(val), 9.5)
    else:
        doc.add_paragraph("本次未识别出风险项。注意：这不等于不存在风险，仅表示在现有材料与核查要点范围内未发现。")

    doc.add_heading("五、逐项说明与整改建议", level=1)
    for idx, f in enumerate(findings, start=1):
        doc.add_heading(f"{idx}. [{f['code']}] {f['title']}（{LEVEL_CN.get(f['risk_level'])}风险）", level=2)
        doc.add_paragraph(f"事实：{f['fact']}")
        doc.add_paragraph(f"依据：{f.get('rule_id') or '—'} {f.get('rule_ref') or '（规范依据待补，见文末说明）'}")
        bound = rule_refs.bound_clauses(f.get("code", ""))
        if bound:
            doc.add_paragraph(
                "评估落点：" + "；".join(f"第 {b['code']} 条（p{b['page']}）" for b in bound)
                + "——该事项是否构成上市障碍，需结合重大性与整改情况判断"
            )
        similar = rule_refs.lookup_similar(f"{f.get('title', '')} {f.get('fact', '')}", limit=2)
        if similar:
            refs = []
            for item in similar:
                parts = [f"「{item['topic']}」{item['note']}"]
                if item["cases"]:
                    parts.append("相似案例：" + "、".join(f"负面清单 #{c}" for c in item["cases"]))
                if item["clause"]:
                    parts.append(f"相关条款：{item['clause']}")
                refs.append("；".join(parts))
            doc.add_paragraph("参考（非依据）：" + " | ".join(refs))
        doc.add_paragraph(f"整改建议：{f['advice']}")
        doc.add_paragraph("证据：")
        any_cite = False
        for c in f["citations"]:
            if not c.get("document_id"):
                continue
            any_cite = True
            doc.add_paragraph(f"《{c['filename']}》第 {c['page']} 页：「{c['quote'][:300]}」", style="List Bullet")
        if not any_cite:
            doc.add_paragraph("（无原文锚点：对应材料缺失，本条仅作补料提示）", style="List Bullet")
        doc.add_paragraph(f"质检：{QC_CN.get(f.get('qc', ''), '未质检')}")

    doc.add_heading("六、质检结论", level=1)
    doc.add_paragraph(
        f"共核查 {qc_result.get('checked', 0)} 条结论；阻断项 {qc_result.get('failed', 0)} 条；"
        f"结果：{'通过' if qc_result.get('passed') else '未通过'}。"
    )
    for issue in qc_result.get("issues", []):
        doc.add_paragraph(f"[{issue['code']}] {issue['title']}：{issue['reason']}", style="List Bullet")
    for notice in qc_result.get("notices", []):
        doc.add_paragraph(f"[{notice['code']}] {notice['title']}：{notice['reason']}", style="List Bullet")
    if qc_result.get("note"):
        doc.add_paragraph(f"注：{qc_result['note']}。")

    if triggers:
        doc.add_heading("七、须律师复核的事项", level=1)
        doc.add_paragraph(f"本次核查触发 {len(triggers)} 项复核条件，律师确认前本报告不得作为最终结论。" )
        for t in triggers:
            doc.add_heading(f"{t['code']} {t['title']}（{t['category']}）", level=2)
            doc.add_paragraph(t["detail"])
            for ev in t.get("evidence", []):
                doc.add_paragraph(ev, style="List Bullet")
        doc.add_heading("八、声明与限制", level=1)
    else:
        doc.add_heading("七、声明与限制", level=1)
    for line in [
        "1. 本报告为机器辅助整理结果，不构成法律意见，不替代律师专业判断。",
        "2. 结论仅基于本次上传材料及其可提取文本；扫描件、缺页、未提供材料的核查项未完成。",
        "3. 每条结论均附证据定位（文件、页码、原文片段）；未附定位的结论视为未验证。",
        "4. 核查要点来源于团队办案经验；规范依据尚在核实中，未核实前不得据以对外作出法律结论。",
        "5. 用户上传材料及其内容不向第三方公开。",
    ]:
        doc.add_paragraph(line)

    # 附：本次报告引用到的上市规则条文原文（只列索引里真实存在的）
    codes: list[str] = []
    for f in findings:
        for clause in rule_refs.bound_clauses(f.get("code", "")):
            if clause["code"] not in codes:
                codes.append(clause["code"])
    if codes:
        doc.add_heading("附：所引上市规则条文", level=1)
        doc.add_paragraph("以下条文摘自港交所《主板上市规则》综合版，页码为 PDF 页序，供核对原文用。")
        for code in codes:
            clause = rule_refs.by_code(code)
            if not clause:
                continue
            doc.add_paragraph(f"第 {clause['code']} 条（p{clause['page']}）", style="List Bullet")
            doc.add_paragraph(f"{clause['text']}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
