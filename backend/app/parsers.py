"""文档解析：PDF / DOCX / TXT / 图片 → 带页码标记的纯文本。

页码标记格式 [[PAGE n]]，供抽取与质检定位证据。
纯图片的 PDF（扫描件）与图片文件走本地 OCR（rapidocr-onnxruntime，不联网）；
OCR 也拿不到文字时标记 needs_ocr，不臆造内容。
"""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

from .config import OCR_DPI, OCR_ENABLED, OCR_MAX_PAGES, OCR_MIN_SCORE

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
_OCR_ENGINE = None


def _clean(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r\n", "\n").split("\n")]
    return "\n".join(lines).strip()


# ------------------------------------------------------------------ OCR
def ocr_ready() -> bool:
    """OCR 是否可用（引擎能导入且未被环境变量关掉）。"""
    if not OCR_ENABLED:
        return False
    try:
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def _engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR

        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _ocr_image(image_path: str) -> str:
    engine = _engine()
    result, _elapse = engine(image_path)
    if not result:
        return ""
    lines: list[str] = []
    for item in result:
        try:
            text, score = item[1], item[2]
        except (IndexError, TypeError):
            continue
        if score is not None and float(score) < OCR_MIN_SCORE:
            continue
        text = (text or "").strip()
        if text:
            lines.append(text)
    return _clean("\n".join(lines))


def _ocr_pdf(path: Path) -> tuple[str, int, str]:
    try:
        import pymupdf as fitz
    except ImportError:  # 旧版 PyMuPDF 只提供 fitz
        import fitz

    doc = fitz.open(str(path))
    total = doc.page_count
    chunks: list[str] = []
    for idx in range(min(total, OCR_MAX_PAGES)):
        page = doc.load_page(idx)
        pix = page.get_pixmap(dpi=OCR_DPI)
        fd, tmp = tempfile.mkstemp(suffix=".png")
        os.close(fd)
        try:
            pix.save(tmp)
            text = _ocr_image(tmp)
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass
        chunks.append(f"[[PAGE {idx + 1}]]\n{text}")
    body = "\n\n".join(chunks)
    ok = len(_clean(body).replace("[[PAGE", "")) > 40
    return body, total, ("ocr" if ok else "needs_ocr")


# ------------------------------------------------------------------ 各格式
def parse_pdf(path: Path) -> tuple[str, int, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    chunks: list[str] = []
    for idx, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        chunks.append(f"[[PAGE {idx}]]\n{_clean(text)}")
    body = "\n\n".join(chunks)
    has_text = len(_clean(body).replace("[[PAGE", "")) > 40
    if has_text:
        return body, len(reader.pages), "parsed"

    if ocr_ready():
        try:
            return _ocr_pdf(path)
        except Exception:  # noqa: BLE001  OCR 失败就当没识别出来，不阻断上传
            pass
    return body, len(reader.pages), "needs_ocr"


def parse_image(path: Path) -> tuple[str, int, str]:
    if not ocr_ready():
        return "", 1, "needs_ocr"
    try:
        text = _ocr_image(str(path))
    except Exception:  # noqa: BLE001
        return "", 1, "failed"
    ok = len(text) > 10
    return f"[[PAGE 1]]\n{text}", 1, ("ocr" if ok else "needs_ocr")


def parse_docx(path: Path) -> tuple[str, int, str]:
    import docx

    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    body = "[[PAGE 1]]\n" + _clean("\n".join(parts))
    ok = len(_clean("\n".join(parts))) > 10
    return body, 1, ("parsed" if ok else "needs_ocr")


def parse_txt(path: Path) -> tuple[str, int, str]:
    for encoding in ("utf-8", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return "", 0, "failed"
    return "[[PAGE 1]]\n" + _clean(text), 1, "parsed"


def parse_file(path: Path) -> tuple[str, int, str]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return parse_pdf(path)
        if suffix == ".docx":
            return parse_docx(path)
        if suffix in IMAGE_SUFFIXES:
            return parse_image(path)
        if suffix in {".txt", ".md", ".csv"}:
            return parse_txt(path)
        return "", 0, "unsupported"
    except Exception:
        return "", 0, "failed"


def page_of_offset(text: str, offset: int) -> int | None:
    """由字符偏移反推页码。"""
    page = None
    for idx, marker in enumerate(_markers(text)):
        if marker > offset:
            break
        page = idx + 1
    return page


def _markers(text: str) -> list[int]:
    import re

    return [m.start() for m in re.finditer(r"\[\[PAGE \d+\]\]", text)]
